import streamlit as st
import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Mm, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
import io
import re

# Configure the Streamlit page and introduce the CSV-to-PowerPoint workflow.
st.set_page_config(page_title="Flashcard PPT Generator", page_icon="📚", layout="centered")

st.title("📚 Flashcard PowerPoint Generator")
st.markdown("Upload your CSV question bank to automatically generate a fresh PowerPoint flashcard deck.")

# Keep the loaded question bank between Streamlit reruns caused by user input.
if "df" not in st.session_state:
    st.session_state.df = None


def clean_value(value):
    """Convert a CSV cell to text suitable for a PowerPoint text frame."""
    if value is None:
        return ""
    # Pandas represents missing numeric cells as NaN; do not print "nan" on a card.
    if isinstance(value, float) and pd.isna(value):
        return ""
    return str(value)


def add_line_breaks_before_numbers(text):
    """Separate inline numbered items onto their own paragraphs for card readability."""
    if text is None:
        return ""
    text = str(text)
    # Split before a number like "2." when it follows other text on the same line.
    return re.sub(r'([^\n])\s+(\d+\.)', r'\1\n\n\2', text)


def create_flashcard_grid(slide, slide_width, slide_height, topics, questions=None, answers=None, is_answer_sheet=False):
    """Place up to eight corresponding question or answer cards on one slide."""
    # Portrait A4 is split into two columns and four rows, in reading order.
    rows = 4
    cols = 2

    # These are the outside page margin and the clear space between neighboring cards.
    margin = Mm(5)
    gap = Mm(3)

    # Divide the remaining page area evenly so front and reverse card edges coincide.
    usable_width = slide_width - (margin * 2) - (gap * (cols - 1))
    usable_height = slide_height - (margin * 2) - (gap * (rows - 1))
    card_width = usable_width / cols
    card_height = usable_height / rows

    # Iterate over every grid position; unused positions on the last page stay blank.
    for row in range(rows):
        for col in range(cols):
            slot_index = row * cols + col
            if is_answer_sheet:
                # Duplex printing on the long edge reverses left/right, so mirror the
                # source card within each row to keep each answer behind its question.
                source_index = row * cols + (cols - 1 - col)
                if source_index >= len(topics):
                    continue
                # The answer side omits the topic/title and places only the answer text.
                title_value = topics[source_index]
                content_value = answers[source_index] if answers is not None and source_index < len(answers) else ""
            else:
                if slot_index >= len(topics):
                    continue
                # The front pairs the topic title with its question text.
                title_value = topics[slot_index]
                content_value = questions[slot_index] if questions is not None and slot_index < len(questions) else ""

            # Each matching front/back position has identical outer geometry.
            left = margin + col * (card_width + gap)
            top = margin + row * (card_height + gap)
            create_single_flashcard(
                slide=slide,
                left=left,
                top=top,
                width=card_width,
                height=card_height,
                title=clean_value(title_value),
                content=clean_value(content_value),
                is_answer_sheet=is_answer_sheet,
            )


def create_single_flashcard(slide, left, top, width, height, title, content, is_answer_sheet=False):
    """Draw a question-side card or its full-size, borderless answer reverse."""
    if is_answer_sheet:
        # The reverse answer area spans the whole card footprint, including the area
        # occupied by the title bar and question body on the front. No reverse border
        # is drawn, avoiding visible offset edges if duplex registration is imperfect.
        content_left = left
        content_top = top
        content_width = width
        content_height = height
    else:
        # Front-side outer card: white paper area with a fine black cut/guide border.
        outer_card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        outer_card.fill.solid()
        outer_card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        outer_card.line.color.rgb = RGBColor(0, 0, 0)
        outer_card.line.width = Pt(0.5)

        # Front-side title block occupies the top 24% of the card height.
        title_area_height = height * 0.24
        title_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, title_area_height)
        title_box.fill.solid()
        title_box.fill.fore_color.rgb = RGBColor(0, 0, 0)
        title_box.line.color.rgb = RGBColor(0, 0, 0)
        title_box.line.width = Pt(0.5)

        # Title text formatting: vertically centered, left aligned, and inset from
        # the title block's left/right edges so the white lettering does not touch them.
        title_tf = title_box.text_frame
        title_tf.clear()
        title_tf.word_wrap = True
        title_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        title_tf.margin_left = Mm(3)
        title_tf.margin_right = Mm(3)
        title_tf.margin_top = Mm(1.5)
        title_tf.margin_bottom = Mm(1.5)
        p = title_tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for run in p.runs:
            # The larger bold white font is specifically for the black title block.
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

        # Front-side question area starts below the title block. The text box itself
        # is inset 2.5 mm on each side and leaves 1.2 mm below the title and 2.4 mm
        # of total vertical clearance; these are external box offsets, not text margins.
        content_left = left + Mm(2.5)
        content_top = top + title_area_height + Mm(1.2)
        content_width = width - Mm(5)
        content_height = height - title_area_height - Mm(2.4)

    # This is the question text box on the front, or the answer text box on the reverse.
    content_box = slide.shapes.add_textbox(content_left, content_top, content_width, content_height)
    content_box.fill.background()
    # Hide the text-box outline: the front card already has its outer border, while
    # the answer reverse intentionally has no border at all.
    content_box.line.color.rgb = RGBColor(255, 255, 255)
    content_box.line.width = Pt(0)

    # Apply the same readability treatment to both the question and answer content.
    body = add_line_breaks_before_numbers(content)
    body_tf = content_box.text_frame
    body_tf.clear()
    body_tf.word_wrap = True
    # Center question/answer paragraphs vertically within their respective text areas.
    body_tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    # The reverse text box is full-card size, so inset answer text 8 mm from its
    # left and right edges. Front question text already uses an inset text box.
    body_tf.margin_left = Mm(8) if is_answer_sheet else Mm(0)
    body_tf.margin_right = Mm(8) if is_answer_sheet else Mm(0)
    # Keep the paragraph area flush vertically; the front question box already has
    # its top/bottom spacing in its position and height above.
    body_tf.margin_top = Mm(0)
    body_tf.margin_bottom = Mm(0)

    if body.strip():
        lines = body.splitlines()
        for idx, line in enumerate(lines):
            # Preserve explicit line breaks as separate PowerPoint paragraphs.
            p = body_tf.paragraphs[0] if idx == 0 else body_tf.add_paragraph()
            p.text = line
            p.alignment = PP_ALIGN.LEFT
            p.space_before = Pt(0)
            p.space_after = Pt(0)
            for run in p.runs:
                # Question and answer body text uses regular black 12 pt type.
                run.font.size = Pt(12)
                run.font.bold = False
                run.font.color.rgb = RGBColor(0, 0, 0)
    else:
        p = body_tf.paragraphs[0]
        p.text = ""


def create_flashcard_presentation(topics, questions, answers, show_answers=True):
    """Create paired portrait A4 question and answer slides for duplex printing."""
    prs = Presentation()
    # Portrait A4 page dimensions, expressed in inches for python-pptx.
    prs.slide_width = Inches(8.27)
    prs.slide_height = Inches(11.69)
    blank_layout = prs.slide_layouts[6]

    num_cards = len(topics)
    # Each question slide is immediately followed by its matching answer reverse.
    cards_per_slide = 8
    num_pairs = (num_cards + cards_per_slide - 1) // cards_per_slide

    for pair_index in range(num_pairs):
        # Slice each column identically to keep topic, question, and answer rows paired.
        start = pair_index * cards_per_slide
        end = min(start + cards_per_slide, num_cards)

        slide_topics = topics[start:end]
        slide_questions = questions[start:end]
        slide_answers = answers[start:end]

        # Front side: show the topic in the title block and the question below it.
        question_slide = prs.slides.add_slide(blank_layout)
        create_flashcard_grid(
            slide=question_slide,
            slide_width=prs.slide_width,
            slide_height=prs.slide_height,
            topics=slide_topics,
            questions=slide_questions,
            answers=None,
            is_answer_sheet=False,
        )

        # Reverse side: show answers in mirrored positions; optionally leave them blank.
        answer_slide = prs.slides.add_slide(blank_layout)
        answer_values = slide_answers if show_answers else ["" for _ in slide_answers]
        create_flashcard_grid(
            slide=answer_slide,
            slide_width=prs.slide_width,
            slide_height=prs.slide_height,
            topics=slide_topics,
            questions=None,
            answers=answer_values,
            is_answer_sheet=True,
        )

    return prs


def parse_csv_flexible(csv_text):
    """Parse pasted Topic,Question,Answer rows, allowing commas in quoted answers."""
    import csv

    # Accept pasted data with or without a surrounding Markdown code fence.
    csv_text = csv_text.strip()
    if csv_text.startswith("```csv"):
        csv_text = csv_text[6:]
    if csv_text.startswith("```"):
        csv_text = csv_text[3:]
    if csv_text.endswith("```"):
        csv_text = csv_text[:-3]

    # Ignore blank lines and remove the optional column header before CSV parsing.
    lines = [line.strip() for line in csv_text.strip().splitlines() if line.strip()]
    if not lines:
        raise ValueError("No data was pasted.")
    if lines[0].lower().startswith("topic,") or lines[0].startswith('"Topic"'):
        lines = lines[1:]

    rows = []
    reader = csv.reader(lines)

    # Require at least three columns and retain Topic, Question, and Answer in order.
    for idx, row in enumerate(reader, start=1):
        if len(row) < 3:
            raise ValueError(
                f"Row {idx} is not in a valid Topic,Question,Answer format. "
                f"Found {len(row)} parts: {row}"
            )
        topic = row[0].strip().strip('"')
        question = row[1].strip().strip('"')
        answer = row[2].strip().strip('"')
        rows.append({"Topic": topic, "Question": question, "Answer": answer})

    return pd.DataFrame(rows)


# Sidebar controls for whether the reverse side prints answers.
st.sidebar.header("Configuration")
show_answers = st.sidebar.checkbox("Include Answers on Answer Slides", value=True, help="Uncheck to generate blank answer slides.")

# Choose between uploading a CSV file and pasting CSV text directly.
st.sidebar.header("Data Input")
input_method = st.sidebar.radio("How would you like to input your data?", ["Upload CSV File", "Paste CSV Data"])

if input_method == "Upload CSV File":
    # Read the uploaded question bank into the shared dataframe session state.
    uploaded_csv = st.sidebar.file_uploader("Upload CSV Question Bank", type=["csv"])
    if uploaded_csv is not None:
        try:
            st.session_state.df = pd.read_csv(uploaded_csv)
        except Exception as e:
            st.error(f"Error reading CSV file: {e}")
else:
    # A form prevents the app from reparsing pasted text on every rerun.
    with st.sidebar.form("csv_input_form"):
        csv_text = st.text_area(
            "Paste your data here in this format: Topic,Question,Answer",
            height=200,
            placeholder="Topic,Question,Answer\nDensity,1. What is density?,1. Density = mass / volume, ρ = m / V [cite: 1]"
        )
        submitted = st.form_submit_button("📤 Parse CSV Data")

    if submitted:
        if csv_text.strip():
            # Parse the pasted CSV and report malformed rows to the user.
            try:
                st.session_state.df = parse_csv_flexible(csv_text)
            except Exception as e:
                st.error(f"Error parsing CSV data: {e}")
        else:
            st.warning("Please paste CSV data before parsing.")


# Preview loaded rows, then use the first three columns as topic/question/answer.
if st.session_state.df is not None:
    try:
        st.success(f"Successfully loaded CSV with {len(st.session_state.df)} rows.")
        st.subheader("CSV Data Preview")
        st.dataframe(st.session_state.df.head())

        if len(st.session_state.df.columns) < 3:
            st.error("The uploaded CSV must have at least 3 columns: Topic, Question, and Answer.")
        else:
            t_col = st.session_state.df.columns[0]
            q_col = st.session_state.df.columns[1]
            a_col = st.session_state.df.columns[2]

            topics = st.session_state.df[t_col].astype(str).tolist()
            questions = st.session_state.df[q_col].astype(str).tolist()
            answers = st.session_state.df[a_col].astype(str).tolist()

            # Generate paired slides in memory and offer the PowerPoint as a download.
            if st.button("Generate Flashcard Presentation"):
                try:
                    prs = create_flashcard_presentation(topics, questions, answers, show_answers=show_answers)

                    output_buffer = io.BytesIO()
                    prs.save(output_buffer)
                    output_buffer.seek(0)

                    # Report the number of question/answer pairs, not individual slides.
                    num_cards = len(questions)
                    num_pairs = (num_cards + 7) // 8
                    st.success(f"Generated presentation with {len(prs.slides)} slides ({num_pairs} question/answer page sets)!")

                    st.download_button(
                        label="📥 Download Generated Flashcards (.pptx)",
                        data=output_buffer,
                        file_name="generated_flashcards.pptx",
                        mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    )
                except Exception as e:
                    st.error(f"Error generating PowerPoint: {e}")
                    import traceback
                    st.error(traceback.format_exc())

    except Exception as e:
        st.error(f"Error processing CSV data: {e}")
else:
    st.info("Please upload a CSV file or paste CSV data via the sidebar to get started.")
